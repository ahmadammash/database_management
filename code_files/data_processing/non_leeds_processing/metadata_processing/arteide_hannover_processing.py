# -----------------------------------------------------------------------------
# Setup & Imports
# -----------------------------------------------------------------------------

# Import the required libraries
import sys
from pathlib import Path

import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# Define the current directory and project root
current_dir = Path(__file__).resolve()
project_root = next(
    (p for p in Path(".").resolve().parents if (p / "pyproject.toml").exists()),
    Path(".").resolve(),
)
if str(project_root) not in sys.path:
    sys.path.insert(0, str(project_root))

# add the root path to the system
if str(project_root) not in sys.path:
    sys.path.insert(0, str(project_root))

# Import the required functions
from code_files.utilities.io.database import create_gcloud_engine

# Change the column and row display settings to show the entire dataframe
pd.set_option("display.max_columns", None)
pd.set_option("display.max_rows", None)

# Define the env_path and create the connection engine
env_path = project_root / "code_files" / "utilities" / "io" / ".env"
engine = create_gcloud_engine(path=env_path)

# Define the base directory
BASE_DIR = Path("C:/Users/medaamm/Box/Data file (AA)/Arteide_Hannover/Source Files")

# -----------------------------------------------------------------------------
# Column Name Renaming Dictionary
# -----------------------------------------------------------------------------

GLOBAL_COLUMN_MAP = {
    "AGE": "Age",
    "AJCC Stage Grouping\n(or other applicable staging system)": "AJCC_Stage_Or_Other_Applicable_System",
    "AJCC Stage Grouping\n(or other applicable staging system) - Ann Arbor": "AJCC_Stage_Ann_Arbor",
    "AJCC Stage Grouping\n(or other applicable staging system) - Ann Arbor Staging System": "AJCC_Stage_Ann_Arbor",
    "AJCC TNM \n(or other applicable)": "AJCC_TNM_Or_Other_Applicable_System",
    "AJCC TNM \n(or other applicable) - Clinical": "AJCC_TNM_Clinical_Or_Other_Applicable",
    "COLLECTION PROCEDURE": "Collection_Procedure",
    "Cancer Status": "Cancer_Status",
    "Concurrent Diseases\n": "Concurrent_Diseases",
    "DESCRIPTION OF TUMOR/\nMACROSCOPIC TUMOR TYPE/\nGROSS DESCRIPTION": "DescriptionTumour_MacroscopicTumourType_GrossDescription",
    "DESCRIPTION OF TUMOR/\nMACROSCOPIC TUMOR TYPE/\nGROSS DESCRIPTION - Gross appereance (size, weight, colour): configuration of tumor (s): e.g. ulcerated, solid-nodule, falt, papillary and indurated, white": "DescriptionTumour_MacroscopicTumourType_GrossDescription",
    "Date of Death": "Date_Of_Death",
    "Date of Disease Recurrence": "Date_Of_Disease_Recurrence",
    "Date of Last Contact ": "Date_Of_Last_Contact",
    "FORMAT": "Format",
    "GENDER": "Gender",
    "GENERAL SUBTYPES OF  CANCER": "General_Cancer_Subtypes",
    "GENERAL SUBTYPES OF CANCER": "General_Cancer_Subtypes",
    "GENERAL SUBTYPES OF MALIGNANT NEOPLASM OF BREAST - ICD CODE": "General_Subtypes_Malignant_Neoplasm_Breast_ICD_Code",
    "GENERAL SUBTYPES OF MALIGNANT NEOPLASMS OF LIVER AND INTRAHEPATIC BILE DUCTS - ICD CODE": "General_Subtypes_Malignant_Neoplasm_Liver_IntrahepticBileDucts_ICD_Code",
    "GENERAL SUBTYPES OF NEOPLASMS OF GALLBLADDER - ICD CODE": "General_Subtypes_Malignant_Neoplasm_Gallbladder_ICD_Code",
    "GENERAL SUBTYPES OF NEOPLASMS OF MALE GENITAL ORGANS - ICD CODE": "General_Subtypes_Malignant_Neoplasm_MaleGenitalOrgans_ICD_Code",
    "GENERAL SUBTYPES OF NEOPLASMS OF MESOTHELIAL AND SOFT TISSUE - ICD CODE": "General_Subtypes_Malignant_Neoplasm_MesothelialAndSoftTissue_ICD_Code",
    "GENERAL SUBTYPES OF NEOPLASMS OF PANCREAS - ICD CODE": "General_Subtypes_Malignant_Neoplasm_Pancreas_ICD_Code",
    "GENERAL SUBTYPES OF NEOPLASMS, STATED OR PRESUMED TO BE PRIMARY, OF LYMPHOID, HEMATOPOIETIC AND RELATED TISSUE - ICD CODE": "General_Subtypes_Neoplasms_Primary_Lymphoid_Hematopoietic_And_Related_ICD_Code",
    "GROSS DESCRIPTION": "Gross_Description",
    "HER2 Status\n - Human Epidermal growth factor Receptor 2 - HER2\n": "HER2_Status_HumanEpidermalGrowthFactorReceptor2",
    "HER2 Status\n - Other ancillary findings": "HER2_Status_Other_Ancillary_Findings",
    "HISTOLOGIC FEATURES/\nMICROSCOPIC DESCRIPTION": "HistologicFeatures_MicroscopicDescription",
    "HISTOLOGIC FEATURES/\nMICROSCOPIC DESCRIPTION - How the cells appear under a microscope: cell struture, tumor margins, vascular invasion, depth of invasion, pathologic stage, growth pattern, Crush Artifact, necrosis, azzopardi effect, apoptosis": "HistologicFeatures_MicroscopicDescription",
    "HISTOLOGIC FEATURES/\nMICROSCOPIC DESCRIPTION - How the cells appear under a microscope: cell struture, tumor margins, vascular invasion, depth of invasion, pathologic stage; growth pattern, Crush Artifact, necrosis, azzopardi effect, apoptosis": "HistologicFeatures_MicroscopicDescription",
    "HISTOLOGICAL GRADE": "Histological_Grade",
    "HISTOLOGICAL TUMOR TYPE": "Histological_Tumour_Type",
    "HISTOLOGICAL TUMOR TYPE\n": "Histological_Tumour_Type",
    "HISTOLOGICAL TUMOR TYPE\n(Lauren)": "Histological_Tumour_Type_Lauren",
    "HISTOLOGICAL TUMOR TYPE\n(WHO)": "Histological_Tumour_Type_WHO",
    "Hormon Receptor Status - Estrogen Receptor - ER": "HormoneReceptorStatus_EstrogenReceptor",
    "Hormon Receptor Status - Immunohistochemical markers": "HormoneReceptorStatus_ImmunohistochemicalMarkers",
    "ID FORMAT FIXATIVE": "ID_Format",
    "Immunhistochemistry": "Immunohistochemistry",
    "Immunhistochemistry - Ki-67 expression (immunohistochemistry):\nHotspot percentage of positive tumor cell nuclei": "Immunohistochemistry_KI67_Expression",
    "Immunohistochemical markers": "Immunohistochemical_Markers",
    "LINE SAMPLE-LIST": "Line_Sample_List",
    "LINE SAMPLE-MANIFEST": "Line_Sample_Manifest",
    "LINE SHIPPING LIST": "Line_Shipping_List",
    "LYMPH NODES\n(Regional Lymph Notes) - Status": "Regional_LymphNodes_Status",
    "Lines of Therapy": "Lines_Of_Therapy",
    "MICRO INVASION - OTHER ORGANS/INVOLVEMENT OTHER ORGANS": "Micro_Invasion_Other_Organs",
    "MICROSCOPIC DESCRIPTION/HISTOLOGIC FEATURES": "Microscopic_Description_Histologic_Features",
    "MONTH-YEAR OF INITIAL DIAGNOSIS OF PRIMARY CANCER": "Month_Year_Initial_Diagnosis_PrimaryCancer",
    "MONTH-YEAR OF INITIAL DIAGNOSIS OF PRIMARY CANCER - DD – MM – YYYY": "Month_Year_Initial_Diagnosis_PrimaryCancer",
    "MONTH-YEAR OF SURGERY-SAMPLE COLLECTION": "Month_Year_Surgery_SampleCollection",
    "MONTH-YEAR OF SURGERY-SAMPLE COLLECTION - DD – MM – YYYY": "Month_Year_Surgery_SampleCollection",
    "MONTH/YEAR OF INITIAL DIAGNOSIS OF PRIMARY CANCER": "Month_Year_Initial_Diagnosis_PrimaryCancer",
    "MONTH/YEAR OF SURGERY/SAMPLE COLLECTION": "Month_Year_Surgery_SampleCollection",
    "Medications\n": "Medications",
    "Metastatic Site\n(Distant)": "Metastatic_Site_Distant",
    "Metastatic Site\n(Distant) - Status": "Metastatic_Site_Distant_Status",
    "Molecular data": "Molecular_Data",
    "NEW DIAGNOSIS": "New_Diagnosis",
    "NEW PRIMARY OR RECURRENCE/\nSPECIMEN NATURE": "New_Primary_Or_Recurrence_SpecimenNature",
    "NEW PRIMARY OR RECURRENCE/\nSPECIMEN NATURE - The nature of the specimen taken during a Clinical Investigation": "New_Primary_Or_Recurrence_SpecimenNature",
    "Other ancillary findings": "Other_Ancillary_Findings",
    "PATIENT-ID": "Patient_ID",
    "POSITION IN BOX": "Position_In_Box",
    "PRIMARY TUMOR SITE": "Primary_Tumour_Site",
    "PROJECT-CODE": "Project_Code",
    "PURCHASE-ORDER": "Purchase_Order",
    "QC-Results - Tumor Tissue %": "QC_Results_TumourTissue_Perc",
    "RACE/\nETHNICITY": "Ethnicity",
    "Recurrent Disease Location (s)": "Recurrent_Disease_Locations",
    "Recurrent Disease Type (s)": "Recurrent_Disease_Types",
    "SAMPLE-FORMAT": "Sample_Format",
    "SAMPLE-ID": "Sample_ID",
    "SAMPLE_FORMAT": "Sample_Format",
    "SITE-ID": "Site_ID",
    "SIZE OF TUMOR": "Tumour_Size",
    "SIZE OF TUMOR\nmm": "Tumour_Size",
    "SIZE OF TUMOR - Tumor Size\n(L x W x D or greatest dimension mm)": "Tumour_Size",
    "SIZE/WEIGHT OF TUMOR\nmm/g": "Tumour_Weight",
    "SPECIMEM NATURE": "Specimen_Nature",
    "SURGICAL MARGIN STATUS": "Surgical_Margin_Status",
    "TUMOR INVASION\nExtent of tumor\n": "TumourInvasion_Extent",
    "TUMOR INVASION - EXTENT OF TUMOR\n": "TumourInvasion_Extent",
    "TUMOR SITE\n(location)": "TumourSite_Location",
    "TYPE OF CANCER-INDICATION\nC50 Malignant neoplasm of breast ": "CancerType_Indication_C50_Malignant_Neoplasm_Of_Breast",
    "TYPE OF CANCER/INDICATION": "CancerType_Indication",
    "TYPES OF CANCER - INDICATION": "CancerType_Indication",
    "TYPES OF CANCER - INDICATION - DIAGNOSIS": "CancerType_Indication_Diagnosis",
    "Tx NAIVE CONFIRMATION?": "Tx_Naive_Confirmation",
    "Unnamed: 20 - Nottingham Histologic Score": "Histological_Grade_Nottingham_Histologic_Score",
    "Unnamed: 21 - Nottingham Histologic Score": "Histological_Grade_Nottingham_Histologic_Score",
    "Unnamed: 23 - No. Assessed": "Lymph_Nodes_Num_Assessed",
    "Unnamed: 24 - No. Assessed": "Lymph_Nodes_Num_Assessed",
    "Unnamed: 24 - No. Positive": "Lymph_Nodes_Num_Positive",
    "Unnamed: 24 - Pathological": "AJCC_TNM_Pathological_Or_Other_Applicable",
    "Unnamed: 25 - AJCC Tumor staging category (T)": "AJCC_pT_Or_Other_Applicable",
    "Unnamed: 25 - No. Assessed": "Lymph_Nodes_Num_Assessed",
    "Unnamed: 25 - No. Positive": "Lymph_Nodes_Num_Positive",
    "Unnamed: 25 - Pathological": "AJCC_TNM_Pathological_Or_Other_Applicable",
    "Unnamed: 25 - Size": "Lymph_Nodes_Size",
    "Unnamed: 26 - AJCC Tumor staging category (T)": "AJCC_pT_Or_Other_Applicable",
    "Unnamed: 26 - AJCC Tumor staging category (pT)": "AJCC_pT_Or_Other_Applicable",
    "Unnamed: 26 - Localization": "Lymph_Nodes_Localization",
    "Unnamed: 26 - No. Assessed": "Lymph_Nodes_Num_Assessed",
    "Unnamed: 26 - No. Positive": "Lymph_Nodes_Num_Positive",
    "Unnamed: 27 - Localization": "Lymph_Nodes_Localization",
    "Unnamed: 27 - No. Assessed": "Lymph_Nodes_Num_Assessed",
    "Unnamed: 27 - No. Positive": "Lymph_Nodes_Num_Positive",
    "Unnamed: 28 - Localization": "Lymph_Nodes_Localization",
    "Unnamed: 28 - No. Assessed": "Lymph_Nodes_Num_Assessed",
    "Unnamed: 28 - No. Positive": "Lymph_Nodes_Num_Positive",
    "Unnamed: 28 - Pathological": "AJCC_TNM_Pathological_Or_Other_Applicable",
    "Unnamed: 29 - AJCC Tumor staging category (pT)": "AJCC_pT_Or_Other_Applicable",
    "Unnamed: 29 - Localization": "Lymph_Nodes_Localization",
    "Unnamed: 29 - No. Assessed": "Lymph_Nodes_Num_Assessed",
    "Unnamed: 29 - No. Positive": "Lymph_Nodes_Num_Positive",
    "Unnamed: 29 - Pathological": "AJCC_TNM_Pathological_Or_Other_Applicable",
    "Unnamed: 29 - Primary Tumor (T)": "AJCC_pT_Or_Other_Applicable",
    "Unnamed: 29 - Primary tumor (T)": "AJCC_pT_Or_Other_Applicable",
    "Unnamed: 30 - AJCC Tumor staging category (pT)": "AJCC_pT_Or_Other_Applicable",
    "Unnamed: 30 - Localization": "Lymph_Nodes_Localization",
    "Unnamed: 30 - No. Positive": "Lymph_Nodes_Num_Positive",
    "Unnamed: 30 - Pathological": "AJCC_TNM_Pathological_Or_Other_Applicable",
    "Unnamed: 30 - Primary Tumor (T)": "AJCC_pT_Or_Other_Applicable",
    "Unnamed: 30 - Primary tumor (T)": "AJCC_pT_Or_Other_Applicable",
    "Unnamed: 31 - AJCC Tumor staging category (pT)": "AJCC_pT_Or_Other_Applicable",
    "Unnamed: 31 - Localization": "Lymph_Nodes_Localization",
    "Unnamed: 31 - Primary Tumor (T)": "AJCC_pT_Or_Other_Applicable",
    "Unnamed: 31 - Primary tumor (T)": "AJCC_pT_Or_Other_Applicable",
    "Unnamed: 31 - Site": "Metastatic_Site",
    "Unnamed: 32 - No. Assessed": "Lymph_Nodes_Num_Assessed",
    "Unnamed: 32 - Pathological": "AJCC_TNM_Pathological_Or_Other_Applicable",
    "Unnamed: 32 - Site": "Metastatic_Site",
    "Unnamed: 33 - AJCC Tumor staging category (pT)": "AJCC_pT_Or_Other_Applicable",
    "Unnamed: 33 - No. Assessed": "Lymph_Nodes_Num_Assessed",
    "Unnamed: 33 - No. Positive": "Lymph_Nodes_Num_Positive",
    "Unnamed: 33 - Site": "Metastatic_Site",
    "Unnamed: 34 - AJCC Tumor staging category (pN)": "AJCC_pN_Or_Other_Applicable",
    "Unnamed: 34 - Localization": "Lymph_Nodes_Localization",
    "Unnamed: 34 - No. Positive": "Lymph_Nodes_Num_Positive",
    "Unnamed: 34 - Site": "Metastatic_Site",
    "Unnamed: 35 - Localization": "Lymph_Nodes_Localization",
    "Unnamed: 36 - Site": "Metastatic_Site",
    "Unnamed: 37 - Site": "Metastatic_Site",
    "Unnamed: 44 - Molecular data": "Molecular_Data",
    "Unnamed: 44 - Progesterone Receptor - PR": "Progestrone_Receptor",
    "Unnamed: 45 - Progesterone Receptor - PR": "Progestrone_Receptor",
    "Unnamed: 6 - Necrotic Tissue %": "QC_Results_Necrotic_Tissue_Perc",
    "Unnamed: 7 - Other Tissue %": "QC_Results_Other_Tissue_Perc",
    "Unnamed: 8 - Sample Surface\nmm x mm\nmm2\nTumour size/total tumour area (mm2)": "QC_Results_Tumour_Size",
    "Vital Status": "Viral_Status",
}


# -----------------------------------------------------------------------------
# Clean, Read & Combine Generic Helper Function
# -----------------------------------------------------------------------------


def load_and_clean_excel(file_path: Path) -> pd.DataFrame:
    """A helper function to read in and clean the datasets"""

    # Define the dictionary for the dataframes
    data_frames = {}

    # Read in the workbook
    xls = pd.ExcelFile(file_path)

    # Define the sheets to exclude
    sheets_excluded = ["Sheet1", "TriStar Specifications-Atreide"]

    # Loop through the workbook sheets
    for sheet_name in xls.sheet_names:
        if sheet_name not in sheets_excluded:
            data_frames[sheet_name] = pd.read_excel(xls, sheet_name, header=1)

    # Define the list of combined columns
    cols_combined = []

    # Loop through the dataframes
    for name, df in data_frames.items():
        # Get the sub-header row
        sub_header = df.iloc[0].fillna("")
        # Combine the header and sub-header row values into one header row
        df.columns = [
            f"{col} - {sub}" if sub != "" else col
            for col, sub in zip(df.columns, sub_header)
        ]

    # Append the column names to the cols_combined list
    for name, df in data_frames.items():
        for col in df.columns:
            cols_combined.append(col)

    # Convert cols_combined to a set to keep the unique column names only
    cols_combined = set(cols_combined)

    # Rename the columns and then remove duplicated or unnamed columns
    for name, df in data_frames.items():
        df = df.rename(columns=GLOBAL_COLUMN_MAP)
        df = df.loc[:, ~df.columns.duplicated()]
        df = df.loc[:, ~df.columns.str.contains("Unnamed")]
        data_frames[name] = df

    # Create the combined hannover dataframe by concatinating all the dataframes
    combined_hannover = pd.concat(data_frames.values(), axis=0, ignore_index=True)

    # Drop rows with no Sample_ID
    combined_hannover = combined_hannover.dropna(subset=["Sample_ID"])

    return combined_hannover


# -----------------------------------------------------------------------------
# Call the function to create the combined dataframe
# -----------------------------------------------------------------------------

combined_hannover = load_and_clean_excel(
    file_path=Path(
        "C:/Users/medaamm/Box/Data file (AA)/Arteide_Hannover/Source Files/Hannover shipments online file 022119.xlsx"
    )
)

# -----------------------------------------------------------------------------
# Write Clean Output to Excel
# -----------------------------------------------------------------------------

excel_output_path = Path(
    "C:/Users/medaamm/OneDrive - University of Leeds/UOL Data Manager/database_management/data_files/non_leeds_data/metadata/arteide_hannover_samples.xlsx"
)
excel_output_path.parent.mkdir(parents=True, exist_ok=True)
sheet_name = "Combined_Hannover"
with pd.ExcelWriter(excel_output_path, engine="openpyxl") as writer:
    combined_hannover.to_excel(writer, sheet_name=sheet_name, index=False)
    worksheet = writer.sheets["Combined_Hannover"]
    width = 25
    # Column widths
    width = 25
    for col_num in range(1, len(combined_hannover.columns) + 1):
        col_letter = get_column_letter(col_num)
        worksheet.column_dimensions[col_letter].width = width

    # Table styling
    max_row, max_col = combined_hannover.shape
    if max_row > 0 and max_col > 0:
        max_col_letter = get_column_letter(max_col)
        tab = Table(
            displayName=sheet_name.replace(" ", "_"),
            ref=f"A1:{max_col_letter}{max_row + 1}",
        )
        style = TableStyleInfo(
            name="TableStyleMedium9",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=True,
        )
        tab.tableStyleInfo = style
        worksheet.add_table(tab)

# -----------------------------------------------------------------------------
# Upload to Google Cloud Engine
# -----------------------------------------------------------------------------

combined_hannover.to_sql(
    name=sheet_name,
    con=engine,
    schema="hannover",
    if_exists="replace",
    index=False,
)
print("Uploaded to GCP successfully.")
