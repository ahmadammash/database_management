from functools import reduce
from pathlib import Path
from openpyxl import load_workbook
import pandas as pd

# Import connection setup from your project directory
from code_files.utilities.io.setup_db_cnxn import setup_cnxn

# 1. Setup paths
project_root, cnxn, cursor, engine, current_dir = setup_cnxn()
env_path = project_root / "code_files" / "utilities" / "io" / ".env"
tma_data_path = project_root / "data_files" / "non_leeds_data" / "tma_data"

tma_files = {
    "TA4546": "TA4546_(Custom Multi Tumour TMA)_with follow up data.xlsx",
    "TA4547": "TA4547_(NSCLC (KRAS mutated & wild type) with histopath data (new version of TA4323))_with follow up data.xlsx",
    "TA4548": "TA4548_(NSCLC (EGFR mutated & wild type) with histopath data (new version of TA4322))_with follow up data.xlsx",
}


def load_excel_tables(filepath: Path) -> dict[str, pd.DataFrame]:
    """Return a dictionary of DataFrames keyed by Excel table name across all worksheets."""
    wb = load_workbook(filepath, data_only=True)
    tables = {}

    for ws in wb.worksheets:
        if not ws.tables:
            continue

        for table_name, table_obj in ws.tables.items():
            # Robustly resolve table coordinate range across openpyxl versions
            if hasattr(table_obj, "ref"):
                table_ref = table_obj.ref
            elif isinstance(table_obj, str):
                if ":" in table_obj:
                    table_ref = table_obj
                else:
                    tbl = ws.tables.get(table_name)
                    table_ref = getattr(tbl, "ref", str(tbl))
            else:
                continue

            try:
                table_data = ws[table_ref]
            except Exception as e:
                print(f"  [Warning] Could not extract range '{table_ref}' for table '{table_name}': {e}")
                continue

            rows = [[cell.value for cell in row] for row in table_data]

            if not rows or len(rows) < 2:
                continue

            raw_headers = rows[0]
            clean_headers = [
                str(h).replace("\xa0", " ").strip().replace("\n", " ")
                if h is not None
                else f"Unnamed_{i}"
                for i, h in enumerate(raw_headers)
            ]

            df = pd.DataFrame(rows[1:], columns=clean_headers)
            df = df.dropna(how="all").reset_index(drop=True)
            tables[table_name] = df

    return tables


def is_data_table(df: pd.DataFrame) -> bool:
    """Filter out non-patient tables, metadata stubs, and tables containing both DonorID and AnonymisedID."""
    cols = [c.strip() for c in df.columns]

    # Exclude metadata key-value reference tables
    if set(cols) == {"Field", "Value"} or cols == ["Field", "Value"]:
        return False

    # Exclude two-column stubs (e.g. ['NO', 'Name'])
    if len(cols) <= 2:
        return False

    # Normalize column names to check for DonorID AND AnonymisedID
    normalized_cols = {
        c.lower().replace("_", "").replace(" ", "") for c in cols
    }
    has_donor_id = any("donorid" in c for c in normalized_cols)
    has_anonymised_id = any("anonymisedid" in c or "anonymizedid" in c for c in normalized_cols)

    # Skip table if it contains both columns
    if has_donor_id and has_anonymised_id:
        return False

    return True


def clean_id_column(series: pd.Series) -> pd.Series:
    """Normalizes identifiers to eliminate integer/float string mismatches."""
    return (
        series.astype(str)
        .str.strip()
        .str.replace(r"\.0$", "", regex=True)
        .replace({"nan": None, "None": None, "": None})
    )


def find_common_key(df_list: list[pd.DataFrame]) -> str:
    """Finds the primary identifier shared across all patient data tables."""
    candidate_keys = [
        "Sample_ID",
        "NO",
        "Original_ID",
        "Patient_ID",
        "Case_ID",
        "DonorID",
        "Donor_Number",
        "Core_ID",
        "NO.",
        "Sample ID",
    ]
    for key in candidate_keys:
        if all(key in df.columns for df in df_list):
            return key

    # Fallback to column intersection if none of the standard keys match
    shared_cols = list(
        reduce(
            lambda acc, d: acc.intersection(set(d.columns)),
            df_list,
            set(df_list[0].columns),
        )
    )
    if shared_cols:
        return shared_cols[0]

    raise KeyError(
        f"No common key found across data tables. Columns present:\n{[list(d.columns) for d in df_list]}"
    )


def merge_two_tables(left: pd.DataFrame, right: pd.DataFrame, join_key: str) -> pd.DataFrame:
    """Outer joins two DataFrames on join_key and coalesces shared columns."""
    l_copy = left.copy()
    r_copy = right.copy()

    # Normalize key values
    l_copy[join_key] = clean_id_column(l_copy[join_key])
    r_copy[join_key] = clean_id_column(r_copy[join_key])

    # Segregate valid from missing IDs to prevent cross-joining nulls
    l_valid = l_copy[l_copy[join_key].notna()]
    r_valid = r_copy[r_copy[join_key].notna()]
    l_nulls = l_copy[l_copy[join_key].isna()]
    r_nulls = r_copy[r_copy[join_key].isna()]

    overlapping_cols = [c for c in l_copy.columns if c in r_copy.columns and c != join_key]

    merged = pd.merge(l_valid, r_valid, on=join_key, how="outer", suffixes=("", "_DROP"))

    # Coalesce overlapping columns into one to prevent _DROP pollution
    for col in overlapping_cols:
        drop_col = f"{col}_DROP"
        if drop_col in merged.columns:
            merged[col] = merged[col].combine_first(merged[drop_col])
            merged.drop(columns=[drop_col], inplace=True)

    if not l_nulls.empty or not r_nulls.empty:
        merged = pd.concat([merged, l_nulls, r_nulls], ignore_index=True)

    return merged


# 2. Main execution loop
combined_tmas = {}

for tma, filename in tma_files.items():
    filepath = tma_data_path / filename
    tables = load_excel_tables(filepath)

    print(f"\n==================== {tma} ====================")

    # Filter out metadata stubs and tables containing DonorID + AnonymisedID
    valid_tables = {}
    for tbl_name, df in tables.items():
        if is_data_table(df):
            valid_tables[tbl_name] = df
            print(f"  [DATA TABLE]    {tbl_name}: {df.shape}")
        else:
            print(f"  [SKIPPED TABLE] {tbl_name}: {list(df.columns[:5])} ... (contains DonorID+AnonymisedID or metadata)")

    table_list = list(valid_tables.values())

    if not table_list:
        print(f"  [Warning] No valid data tables found in {tma}")
        continue

    if len(table_list) == 1:
        combined_tmas[tma] = table_list[0]
    else:
        key = find_common_key(table_list)
        print(f"  -> Merging {len(table_list)} tables using key: '{key}'")
        combined_tmas[tma] = reduce(
            lambda left, right: merge_two_tables(left, right, join_key=key),
            table_list,
        )

    print(f"  -> Combined shape: {combined_tmas[tma].shape}")

# 3. Access individual datasets separately
ta4546_combined = combined_tmas.get("TA4546")
ta4547_combined = combined_tmas.get("TA4547")
ta4548_combined = combined_tmas.get("TA4548")


# In[]:

COLUMN_MAPPING = {
    # --- Identifiers ---
    "Original ID": "Original_ID",
    "SampleID": "Sample_ID",
    "Sample_ID2": "Sample_ID",
    "ID": "Source_Row_ID",
    # --- Demographics ---
    "AgeAtDX": "Age_At_DX",
    "ETHNICITY2": "Ethnicity",
    "GeographicOrigin": "Ethnicity",
    # --- Staging & Subtyping ---
    "StageAtDX": "Stage_At_DX",
    "AJCC_Stage_Or_Other_Applicable_System": "AJCC_Stage",
    "AJCC_TNM_Clinical_Or_Other_Applicable": "TNM_Clinical",
    "AJCC_TNM_Pathological_Or_Other_Applicable": "TNM_Pathological",
    "AJCC_pT_Or_Other_Applicable": "pT_Stage",
    "pTNM_Stage": "pTNM_Stage",
    "Sub-Type": "Sub_Type",
    "Sub-type11": "Sub_Type",
    "Histopathologic_Subtype": "Histological_Tumour_Type",
    # --- Specimen & Tissue Attributes ---
    "Specimen1_Type": "Specimen_Type",
    "Sample_Type": "Specimen_Type",
    "Specimen_1_Collection_Date": "Date_Of_Collection",
    "Collection_Date": "Date_Of_Collection",
    "Specimen_Collection_Date": "Date_Of_Collection",
    "Year of Surgery": "Surgery_Year",
    "SurgeryDate": "Date_Of_Surgery",
    "Month_Year_Surgery_SampleCollection": "Surgery_Collection_MonthYear",
    "Month_Year_Initial_Diagnosis_PrimaryCancer": "Initial_Diagnosis_MonthYear",
    "New_Primary_Or_Recurrence_SpecimenNature": "Specimen_Nature",
    "TumourSite_Location": "Primary_Tumour_Site",
    "Surgery_Site": "Primary_Tumour_Site",
    "Site_Of_Diagnosis": "Primary_Tumour_Site",
    # --- Biomarkers (IHC, FISH, Assays) ---
    "HER2": "HER2_Status",
    "HER22": "HER2_Status",
    "Her26": "HER2_Status",
    "PD-L1CPS": "PDL1_CPS",
    "PDL1_CPS_2": "PDL1_CPS",
    "PDL1_TPS_2": "PDL1_TPS",
    "PDL1_Percent": "PDL1_TPS",
    # --- Systemic & Adjuvant Therapies ---
    "Adjuvant_ChemoTx_Regimen_1": "Adjuvant_ChemoTx_1",
    "Responder_NonResponder": "Responder_Status",
    "BestResponse_ResponseNonResponse": "Best_Response",
    # --- Recurrence, Metastasis & Relapse ---
    "Metastatic_Site_Distant_Status": "Metastatic_Status",
    # --- Survival & Mortality ---
    "Death_Date": "Date_Of_Death",
    "Date_Last_FollowUp": "Date_Last_Follow_Up",
    "Last_FollowUp": "Last_Follow_Up_Status",
    "Last_FollowUp_Date": "Date_Last_Follow_Up",
    "OS": "OS_Months",
    "OverallSurvival_Months": "OS_Months",
    "DFS": "DFS_Months",
    "DiseaseFreeSurvival_Months": "DFS_Months",
    "Survival_Months": "OS_Months",
    # --- Genomic & Variant Call Fields ---
    "proteinChange": "Protein_Change",
    "location9": "Variant_Location",
}


CANONICAL_ORDER = [
    # =========================================================
    # 1. IDENTIFIERS & SOURCE TRACKING
    # =========================================================
    "NO",
    "Patient_ID",
    "Sample_ID",
    "Original_ID",
    "Vendor_Sample_ID",
    "Biobank_ID",
    "Pathology_ID",
    "Lab_ID",
    "Project_Code",
    "Purchase_Order",
    "Site_ID",
    "Source_Row_ID",
    "File_Name",
    "TMA_Number",
    "Block",
    "Block_A",
    "Block_B",
    "Block_C",
    "Block_D",
    "Block_E",
    # =========================================================
    # 2. PATIENT DEMOGRAPHICS & CLINICAL BASELINE
    # =========================================================
    "Name",
    "Gender",
    "Age",
    "Age_At_DX",
    "Ethnicity",
    "Origin",
    "Concurrent_Medical_Conditions",
    # =========================================================
    # 3. PRIMARY DIAGNOSIS & TOPOGRAPHY
    # =========================================================
    "Diagnosis",
    "Clinical_Diagnosis",
    "Pathological_Diagnosis",
    "Histological_Diagnosis",
    "CancerType_Indication",
    "General_Cancer_Subtypes",
    "Type",
    "Sub_Type",
    "Primary_Tumour_Site",
    "Initial_Diagnosis_MonthYear",
    "Date_Of_Clinical_DX",
    "Date_Of_Histological_Diagnosis",
    # =========================================================
    # 4. HISTOPATHOLOGY, GRADE & STAGING
    # =========================================================
    "Grade",
    "Histological_Grade",
    "Nuclear_Grade",
    "Histological_Tumour_Type",
    "Stage",
    "Stage_At_DX",
    "AJCC_Stage",
    "pTNM_Stage",
    "TNM_Clinical",
    "TNM_Pathological",
    "pT_Stage",
    "T",
    "N",
    "M",
    "Regional_LymphNodes_Status",
    "Lymph_Nodes_Localization",
    "Lymph_Nodes_Num_Assessed",
    "Lymph_Nodes_Num_Positive",
    "TumourInvasion_Extent",
    "Invasive_Edge_Tumor",
    "Gross_Description",
    "HistologicFeatures_MicroscopicDescription",
    # =========================================================
    # 5. BIOSPECIMEN DETAILS & COLLECTION
    # =========================================================
    "Specimen_Type",
    "Specimen_Nature",
    "Sample_Format",
    "Surgical_Specimen",
    "Collection_Procedure",
    "Tumour_Size",
    "Date_Of_Surgery",
    "Surgery_Year",
    "Surgery_Collection_MonthYear",
    "Date_Of_Collection",
    "Collection_Year",
    # =========================================================
    # 6. BIOMARKERS, IHC & TARGETED ASSAYS
    # =========================================================
    # Standard Breast / GYN / General IHC
    "ER",
    "PR",
    "P",
    "HER2_Status",
    "HER2_Score",
    "HER2_ISH",
    "HER3",
    # Immune Checkpoints & Tumor Infiltrating Lymphocytes
    "PDL1_TPS",
    "PDL1_CPS",
    "CD3_per_sqmm",
    "CD8_per_sqmm",
    # Mismatch Repair / MSI
    "MLH1",
    "MSH2",
    "MSH6",
    "PMS1",
    "MSI_Status",
    # Quality Control & Transcript Markers
    "PPIB",
    "RNAscope",
    "Bl",
    # ADC Targets: B7-H3
    "B7H3_IHCResult",
    "B7H3_HScore",
    "B7H3_NegPerc",
    "B7H3_1Plus_2Plus_3Plus_Perc",
    "B7H3_2Plus_3Plus_Perc",
    "B7H3_3Plus_Perc",
    # ADC Targets: B7-H4
    "B7H4_IHCResult",
    "B7H4_HScore",
    "B7H4_NegPerc",
    "B7H4_1Plus_2Plus_3Plus_Perc",
    "B7H4_2Plus_3Plus_Perc",
    "B7H4_3Plus_Perc",
    # ADC Targets: Nectin-4
    "Nectin4_IHCResult",
    "Nectin4_HScore",
    "Nectin4_NegPerc",
    "Nectin4_1Plus_2Plus_3Plus_Perc",
    "Nectin4_2Plus_3Plus_Perc",
    "Nectin4_3Plus_Perc",
    # ADC Targets: TROP-2 (1:1000 & 1:3000 Dilutions)
    "TROP2_1_1000_IHCResult",
    "TROP2_1_1000_HScore",
    "TROP2_1_1000_NegPerc",
    "TROP2_1_1000_1Plus_2Plus_3Plus_Perc",
    "TROP2_1_1000_2Plus_3Plus_Perc",
    "TROP2_1_1000_3Plus_Perc",
    "TROP2_1_3000_IHCResult",
    "TROP2_1_3000_HScore",
    "TROP2_1_3000_NegPerc",
    "TROP2_1_3000_1Plus_2Plus_3Plus_Perc",
    "TROP2_1_3000_2Plus_3Plus_Perc",
    "TROP2_1_3000_3Plus_Perc",
    # =========================================================
    # 7. GENOMICS & MOLECULAR ALTERATIONS
    # =========================================================
    "EGFR_Mutation",
    "KRAS_Mutation",
    "BRCA",
    "gene",
    "transcript",
    "var_type",
    "Protein_Change",
    "Variant_Location",
    "function",
    "vcf.rownum",
    "rowtype",
    # =========================================================
    # 8. SYSTEMIC THERAPIES, SURGERY & RADIATION
    # =========================================================
    "Neoadjuvant_Treatment_Type",
    "Neoadjuvant_Treatment_StartDate",
    "Neoadjuvant_Treatment_EndDate",
    "Start_Treatment_1",
    "Therapy_1",
    "Follow_Up_1",
    "Therapy_1_FollowUp",
    "Therapy_2",
    "Follow_Up_2",
    "Therapy_2_FollowUp",
    "Adjuvant_ChemoTx_1",
    "Adjuvant_ChemoTx_1_Start_Date",
    "Adjuvant_ChemoTx_1_End_Date",
    "Metastatic_Chemo_Targeted_Tx_1",
    "Metastatic_Chemo_Targeted_Tx_1_Start_Date",
    "Immuno_Tx",
    "Immuno_Tx_Start_Date",
    "Radiation_1_Type",
    "Responder_Status",
    "Best_Response",
    # =========================================================
    # 9. RECURRENCE, PROGRESSION & METASTASIS
    # =========================================================
    "Date_Of_Relapse",
    "Other_Relapse_1",
    "Other_Relapse_1_Sites",
    "Other_Relapse_2",
    "Other_Relapse_2_Sites",
    "Metastatic_Status",
    "Metastatic_Site",
    # =========================================================
    # 10. SURVIVAL, VITAL STATUS & FOLLOW-UP
    # =========================================================
    "Current_Status",
    "Death",
    "Date_Of_Death",
    "Cause_Of_Death",
    "Date_Last_Seen",
    "Date_Last_Follow_Up",
    "OS_Months",
    "PFS_Months",
    "DFS_Months",
    # =========================================================
    # 11. LOGISTICS & ANNOTATIONS
    # =========================================================
    "Position_In_Box",
    "Line_Shipping_List",
    "TRK_Incoming",
    "Divided",
    "Others",
]

def harmonize_and_order_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    """Coalesces duplicated columns, applies harmonization mapping,

    and sorts columns into canonical order.
    """
    df = df.copy()

    # 1. Clean hidden whitespace and non-breaking spaces from column headers
    df.columns = [c.replace("\xa0", " ").strip() for c in df.columns]

    # 2. Coalesce duplicate concepts into standard targets
    for raw_col, canonical_col in COLUMN_MAPPING.items():
        if raw_col in df.columns:
            if canonical_col in df.columns and raw_col != canonical_col:
                df[canonical_col] = df[canonical_col].combine_first(df[raw_col])
                df.drop(columns=[raw_col], inplace=True)
            else:
                df.rename(columns={raw_col: canonical_col}, inplace=True)

    # 3. Sort columns by canonical order, keeping remaining columns at the tail
    present_canonical = [c for c in CANONICAL_ORDER if c in df.columns]
    unmapped_cols = [c for c in df.columns if c not in CANONICAL_ORDER]

    return df[present_canonical + unmapped_cols]


ta4546_harmonized = harmonize_and_order_dataframe(ta4546_combined)
ta4547_harmonized = harmonize_and_order_dataframe(ta4547_combined)
ta4548_harmonized = harmonize_and_order_dataframe(ta4548_combined)