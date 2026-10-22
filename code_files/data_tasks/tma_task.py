# Import the required libraries
import pandas as pd
from openpyxl import load_workbook

pd.set_option("display.max_rows", None)
pd.set_option("display.max_columns", None)

# Load the TMA workbook
tma_wb = load_workbook(
    filename="C:/Users/medaamm/Box/Data file (AA)/Miscellaneous Requests/69574527 NSCLC EGFR KRAS WT (TA4527).xlsx"
)

# Load the data workbook
data_wb = load_workbook(
    filename="C:/Users/medaamm/Box/Data file (AA)/Miscellaneous Requests/SOW2523 NSCLC EGFR, KRAS mutated & WT.xlsx"
)

# Define the set of table names to be selected
tbls_to_select = {
    "tma_data",
    "leeds_data",
    "leeds_supp",
    "nonleeds_data_1",
    "nonleeds_data_2",
}

# Extract the target tables
tables = {}

# Define the list of workbooks
wbs = [tma_wb, data_wb]

for wb in wbs:
    for ws in wb.worksheets:
        for table_name, table in ws.tables.items():
            if table_name in tbls_to_select:
                rows = [[cell.value for cell in row] for row in ws[table]]
                df = pd.DataFrame(rows[1:], columns=rows[0])
                tables[f"{table_name}"] = df

# Merge the leeds follow up and supplementary data
leeds_combined = pd.merge(
    tables["leeds_data"],
    tables["leeds_supp"],
    on="DonorID",
    how="inner",
    suffixes=("", "_new"),
)
for col in ["SelSpecAnonymisedID", "KRAS_Mutation", "EGFR_Mutation"]:
    leeds_combined[col] = leeds_combined[col].combine_first(
        leeds_combined[f"{col}_new"]
    )
    leeds_combined.drop(columns=f"{col}_new", inplace=True)


# Concatenate the combined leeds data and the non-leeds data
data_combined = pd.concat(
    [leeds_combined, tables["nonleeds_data_1"], tables["nonleeds_data_2"]]
).reset_index(drop=True)

# Combine some duplicated columns
data_combined["AgeDX"] = data_combined["AgeDX"].combine_first(data_combined["Age"])
tnm_columns = ["T", "N", "M"]
for col in tnm_columns:
    data_combined[col] = data_combined[col].combine_first(data_combined[f"p{col}"])
data_combined["TNM_Stage"] = data_combined["Pathological_TNM_Stage"].combine_first(
    data_combined["TNM_Stage"]
)
data_combined["StageAtDX"] = data_combined["StageAtDX"].combine_first(
    data_combined["Stage"]
)
data_combined["SpecSurgeryDate"] = data_combined["SpecSurgeryDate"].combine_first(
    data_combined["SurgeryDate_MMYY"]
)
data_combined["SpecTumourSite"] = data_combined["SpecTumourSite"].combine_first(
    data_combined["Diagnosis_Site"]
)
data_combined["SpecBlockNum1"] = (
    data_combined["SpecBlockNum1"]
    .combine_first(data_combined["BI"])
    .combine_first(data_combined["Bl"])
)
data_combined["SelSpecAnonymisedID"] = data_combined[
    "SelSpecAnonymisedID"
].combine_first(data_combined["Sample_ID"])
data_combined["Diagnosis_Combined"] = (
    data_combined["Type"] + "-" + data_combined["Sub-Type"]
)
data_combined["Diagnosis_Combined"] = data_combined["Diagnosis_Combined"].combine_first(
    data_combined["Diagnosis"]
)
data_combined["SpecHistologicalDiagnosis"] = (
    data_combined["SpecHistologicalDiagnosis"]
    .combine_first(data_combined["Histology"])
    .combine_first(data_combined["Histological_Diagnosis"])
    .combine_first(data_combined["Diagnosis_Combined"])
)
data_combined["SpecSurgery"] = data_combined["SpecSurgery"].combine_first(
    data_combined["Surgery_And_Site"]
)
# Drop duplicated columns
cols_to_drop = [
    "Age",
    "pT",
    "pN",
    "pM",
    "Pathological_TNM_Stage",
    "Stage",
    "Histology",
    "Histological_Diagnosis",
    "SurgeryDate_MMYY",
    "Diagnosis_Site",
    "BI",
    "Bl",
    "Specimen_1_CollectionDate_MMYY",
    "Sample_ID",
    "Type",
    "Sub-Type",
    "Diagnosis",
    "Diagnosis_Combined",
]
data_combined.drop(columns=cols_to_drop, inplace=True)
# Format the KRAS_Mutation column
for idx, row in data_combined.iterrows():
    kras = row["KRAS_Mutation"]
    protein = row["Protein_Change"]
    if kras == "KRAS Mutated":
        data_combined.loc[idx, "KRAS_Mutation"] = protein
    elif kras == "KRAS-WT":
        data_combined.loc[idx, "KRAS_Mutation"] = "Not detected"

# reorder the columns logically
ordered_columns = [
    # 1. Patient Identifiers & Governance
    "DonorID",
    "Name",
    "Group",
    "Original_ID",
    "Biobank_ID",
    "Pathology_ID",
    "SelSpecAnonymisedID",
    # 2. Patient Demographics & Baseline
    "Gender",
    "AgeDX",
    "Clinical_Diagnosis_Date",
    "DateDX",
    # 3. Clinical Staging & Histology
    "SpecTumourSite",
    "SpecHistologicalDiagnosis",
    "Histological_Diagnosis_Date",
    "Grade",
    "T",
    "N",
    "M",
    "cM",
    "TNM_Stage",
    "StageAtDX",
    "StagingCriteria",
    "StagingGradingInfoProvidedFor",
    # 4. Primary Specimen & Treatment
    "Collection_Date",
    "SpecSurgeryDate",
    "SpecSurgery",
    "Surgery_And_Site",
    "ResectionMarginsClear",
    "AnyNeoadjuvantTherapy",
    "AnyImmunotherapy",
    "SpecTreatmentSubgrouping",
    "SpecIsRecurrence",
    "SpecBlockNum1",
    "SpecBlockNum2",
    "Specimen_1_Type",
    "SpecNumber_Used_For_MolecularTesting_IfDifferent",
    "Block_Used_For_Testing",
    # 5. Matched Metastatic Samples
    "AnonymisedID_MatchedMet1",
    "MatchedMet1_Site",
    "MatchedMet1_SurgeryDate",
    "MatchedMet1_BlockNumber",
    "MatchedMet1_TreatmentSubgrouping",
    "AnonymisedID_MatchedMet2",
    "MatchedMet2_Site",
    "MatchedMet2_SurgeryDate",
    "MatchedMet2_BlockNumber",
    "MatchedMet2_TreatmentSubgrouping",
    # 6. Molecular, Biomarker & Genomic Profile
    "Detection_Platform",
    "gene",
    "Protein_Change",
    "EGFR_Mutation",
    "KRAS_Mutation",
    "NRAS_Mutation",
    "BRAF_Mutation",
    "ALK_Rearrangement_Fusion",
    "PDL1_PosNeg",
    "Percent_Cells_Pos",
    "Any_Other_Variants_Detected",
    # 7. Metadata / Notes
    "Notes",
]
data_combined = data_combined[ordered_columns]
