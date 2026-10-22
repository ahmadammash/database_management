import openpyxl
import pandas as pd

path_ta433 = "C:/Users/medaamm/Box/Data file (AA)/Miscellaneous Requests/69574509 & 69574333/69574333 NSCLC (ADC & SCC) with low, moderate & high PD-L1(22C3) expression with histopath data - optional follow up data (TA4333).xlsx"
path_ta4509 = "C:/Users/medaamm/Box/Data file (AA)/Miscellaneous Requests/69574509 & 69574333/69574509 NSCLC (ADC & SCC) with histopath data, PD-L1(22C3) status & Oncomine Precision Assay NGS data.xlsx"

wb_ta433 = openpyxl.load_workbook(path_ta433, data_only=True)
wb_ta4509 = openpyxl.load_workbook(path_ta4509, data_only=True)

sheet_ta433 = wb_ta433["TA4333 (Follow up Data)"]
sheet_ta4509 = wb_ta4509["TA4509 (Follow Up Data)"]
sheets = [sheet_ta433, sheet_ta4509]

column_name_mapping = {
    "Bl": "Block_Number",
    "CPS(PDL1)": "CPS_PDL1",
    "Core ID": "Core_ID",
    "HER2 ISH": "HER2_ISH",
    "HER2 score": "HER2_Score",
    "NO": "Core_ID",
    "Original ID": "Original_ID",
    "PD-L1(TPS)": "TPS_PDL1",
    "PD-L1CPS": "CPS_PDL1",
    "SAMPLE ID#": "Sample_ID",
    "Sub-Type": "Sub_Type",
    "TPS(PDL1)": "TPS_PDL1",
    "Date_Of_Collection": "Collection_Date",
    "Diagnosis": "Histological_Diagnosis",
}

cols_order = [
    "Core_ID",
    "Sample_ID",
    "Biobank_ID",
    "Pathology_ID",
    "Original_ID",
    "Origin",
    "TRK_Incoming",
    "Block_Number",
    "Gender",
    "Age",
    "Ethnicity",
    "Age_At_DX",
    "Collection_Date",
    "Specimen1_Type",
    "Specimen_1_Collection_Date",
    "Year_Of_Collection",
    "Date_Of_Histological_Diagnosis",
    "Date_Of_Clinical_DX",
    "Date_Of_Surgery",
    "Site_Of_Diagnosis",
    "Surgery_Site",
    "Histological_Diagnosis",
    "Type",
    "Sub_Type",
    "Grade",
    "Stage",
    "T",
    "N",
    "M",
    "Date_Of_Met_Disease",
    "Met_Disease",
    "Metastatic_Chemo_Targeted_Tx_2",
    "Metastatic_Chemo_Targeted_Tx_2_Start_Date",
    "Metastatic_Chemo_Targeted_Tx_2_End_Date",
    "Start_Treatment_1",
    "Therapy_1",
    "Therapy_2",
    "Adjuvant_ChemoTx_1",
    "Adjuvant_ChemoTx_1_Start_Date",
    "Adjuvant_ChemoTx_1_End_Date",
    "Follow_Up_1",
    "Follow_Up_2",
    "Immuno_Met_Tx",
    "Immuno_Met_Tx_Start_Date",
    "Immuno_Met_Tx_End_Date",
    "Radiation_1_Type",
    "Radiation_1_Start_Date",
    "Radiation_1_End_Date",
    "Other_Relapse_1",
    "Other_Relapse_1_Sites",
    "Other_Relapse_2",
    "Other_Relapse_2_Sites",
    "KRAS_Mutation",
    "EGFR_Mutation",
    "HER2",
    "HER2_ISH",
    "HER2_Score",
    "HER3",
    "CPS_PDL1",
    "TPS_PDL1",
    "B7-H3 - 1+,2+,3+(%)",
    "B7-H3 - 2+,3+(%)",
    "B7-H3 - 3+(%)",
    "B7-H3 - H score",
    "B7-H3 - IHC Result",
    "B7-H3 - Neg(%)",
    "B7-H4 - 1+,2+,3+(%)",
    "B7-H4 - 2+,3+(%)",
    "B7-H4 - 3+(%)",
    "B7-H4 - IHC Result",
    "B7-H4 - Neg(%)",
    "B7-H4 -H score",
    "Nectin-4 - 1+,2+,3+(%)",
    "Nectin-4 - 2+,3+(%)",
    "Nectin-4 - 3+(%)",
    "Nectin-4 - H score",
    "Nectin-4 - IHC Result",
    "Nectin-4 - Neg(%)",
    "TROP2 (1:1000) - 1+,2+,3+(%)",
    "TROP2 (1:1000) - 2+,3+(%)",
    "TROP2 (1:1000) - 3+(%)",
    "TROP2 (1:1000) - H score",
    "TROP2 (1:1000) - IHC Result",
    "TROP2 (1:1000) - Neg(%)",
    "TROP2 (1:3000) - 1+,2+,3+(%)",
    "TROP2 (1:3000) - 2+,3+(%)",
    "TROP2 (1:3000) - 3+(%)",
    "TROP2 (1:3000) - H score",
    "TROP2 (1:3000) - IHC Result",
    "TROP2 (1:3000) - Neg(%)",
    "Date_Last_Follow_Up",
    "Last_Follow_Up",
    "Date_Last_Seen",
    "Current_Status",
    "Concurrent_Medical_Conditions",
    "Death_Date",
    "Cause_Of_Death",
    "OS",
    "DFS",
    "Survival_Months",
]

processed_dfs = {}

for sheet in sheets:
    for table_name in sheet.tables:
        table_obj = sheet.tables[table_name]
        table_range = sheet[table_obj.ref]

        # 1. Rename cells in openpyxl worksheet
        header_row = table_range[0]
        for cell in header_row:
            if cell.value:
                clean_name = str(cell.value).strip()
                cell.value = column_name_mapping.get(clean_name, clean_name)

        # 2. Build index map of renamed headers
        col_map = {
            str(cell.value).strip(): idx
            for idx, cell in enumerate(header_row, start=1)
            if cell.value is not None
        }

        # 3. Compute final target column list
        existing_cols_order = [c for c in cols_order if c in col_map]
        unmapped_cols = [c for c in col_map if c not in existing_cols_order]
        final_col_order = existing_cols_order + unmapped_cols

        # 4. Reorder table data
        new_table_data = []
        for row in table_range:
            new_row = [row[col_map[c] - 1].value for c in final_col_order]
            new_table_data.append(new_row)

        # 5. Overwrite cells in memory
        min_col, min_row, _, _ = openpyxl.utils.range_boundaries(table_obj.ref)
        for r_idx, row_values in enumerate(new_table_data, start=min_row):
            for c_idx, val in enumerate(row_values, start=min_col):
                sheet.cell(row=r_idx, column=c_idx, value=val)

        # 6. Update openpyxl table columns (providing integer ID)
        table_obj.tableColumns.clear()
        for col_idx, c_name in enumerate(final_col_order, start=1):
            table_obj.tableColumns.append(
                openpyxl.worksheet.table.TableColumn(id=col_idx, name=c_name)
            )

        # 7. Convert to DataFrame for in-memory inspection
        df_data = new_table_data[1:]  # Data rows
        df_headers = new_table_data[0]  # Header row
        processed_dfs[table_name] = pd.DataFrame(df_data, columns=df_headers)

# Inspect column order and head of a table without modifying files
for t_name, df in processed_dfs.items():
    print(f"\n--- Table: {t_name} ---")
    print("Columns:", list(df.columns))
    print(df.head(2))
