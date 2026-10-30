from functools import reduce
from pathlib import Path
from openpyxl import load_workbook
import pandas as pd

# Import functions from project directory
from code_files.utilities.io.setup_db_cnxn import setup_cnxn

# 1. Setup paths
project_root, cnxn, cursor, engine, current_dir = setup_cnxn()
env_path = project_root / "code_files" / "utilities" / "io" / ".env"
tma_data_path = project_root / "data_files" / "non_leeds_data" / "tma_data"

tma_files = {
    "TA4546": "TA4546_(Custom Multi Tumour TMA)_with follow up data.xlsx",
    "TA4547": "TA4547_(NSCLC (KRAS mutated & wild type) with histopath data (new version of TA4323))_with follow up data.xlsx",
    "TA4548": "TA4548_(NSCLC (EGFR mutated & wild type) with histopath data (new version of TA4322))_with follow up data.xlsx",
}


def load_excel_tables(filepath: Path) -> dict[str, pd.DataFrame]:
    """Return a dictionary of DataFrames keyed by Excel table name across all worksheets."""
    wb = load_workbook(filepath, data_only=True)
    tables = {}

    for ws in wb.worksheets:
        if not ws.tables:
            continue

        for table_name, table_obj in ws.tables.items():
            # Extract coordinate range regardless of openpyxl representation
            if hasattr(table_obj, "ref"):
                table_ref = table_obj.ref
            elif isinstance(table_obj, str):
                if ":" in table_obj:
                    table_ref = table_obj
                else:
                    tbl = ws.tables.get(table_name)
                    table_ref = getattr(tbl, "ref", str(tbl))
            else:
                continue

            try:
                table_data = ws[table_ref]
            except Exception as e:
                print(
                    f"  [Warning] Could not extract range '{table_ref}' for table '{table_name}': {e}"
                )
                continue

            rows = [[cell.value for cell in row] for row in table_data]

            if not rows or len(rows) < 2:
                continue

            raw_headers = rows[0]
            clean_headers = [
                str(h).replace("\xa0", " ").strip().replace("\n", " ")
                if h is not None
                else f"Unnamed_{i}"
                for i, h in enumerate(raw_headers)
            ]

            df = pd.DataFrame(rows[1:], columns=clean_headers)
            df = df.dropna(how="all").reset_index(drop=True)
            tables[table_name] = df

    return tables


def clean_id_column(series: pd.Series) -> pd.Series:
    """Normalizes ID columns to eliminate float/int string mismatches while preserving real nulls."""
    return (
        series.astype(str)
        .str.strip()
        .str.replace(r"\.0$", "", regex=True)
        .replace({"nan": None, "None": None, "": None})
    )


def find_common_key(df_list: list[pd.DataFrame]) -> str:
    """Finds the shared identifier across all dataframes in a workbook."""
    candidate_keys = [
        "Patient_ID",
        "Case_ID",
        "Sample_ID",
        "Original_ID",
        "DonorID",
        "Donor_Number",
        "Core_ID",
        "NO.",
        "Sample ID",
    ]
    for key in candidate_keys:
        if all(key in df.columns for df in df_list):
            return key

    shared_cols = list(
        reduce(
            lambda acc, d: acc.intersection(set(d.columns)),
            df_list,
            set(df_list[0].columns),
        )
    )
    if shared_cols:
        return shared_cols[0]

    raise KeyError(
        f"No common key found across tables. Columns present: {[list(d.columns) for d in df_list]}"
    )


def merge_two_tables(
    left: pd.DataFrame, right: pd.DataFrame, join_key: str
) -> pd.DataFrame:
    """Safely outer joins two DataFrames on join_key and coalesces shared columns."""
    l_copy = left.copy()
    r_copy = right.copy()

    # Standardize join key format
    l_copy[join_key] = clean_id_column(l_copy[join_key])
    r_copy[join_key] = clean_id_column(r_copy[join_key])

    # Avoid cross-joining on rows where ID is null
    l_valid = l_copy[l_copy[join_key].notna()]
    r_valid = r_copy[r_copy[join_key].notna()]
    l_nulls = l_copy[l_copy[join_key].isna()]
    r_nulls = r_copy[r_copy[join_key].isna()]

    overlapping_cols = [
        c for c in l_copy.columns if c in r_copy.columns and c != join_key
    ]

    merged = pd.merge(
        l_valid, r_valid, on=join_key, how="outer", suffixes=("", "_DROP")
    )

    # Coalesce overlapping columns into one to prevent _DROP pollution
    for col in overlapping_cols:
        drop_col = f"{col}_DROP"
        if drop_col in merged.columns:
            merged[col] = merged[col].combine_first(merged[drop_col])
            merged.drop(columns=[drop_col], inplace=True)

    # Re-append rows that had null keys (if any exist)
    if not l_nulls.empty or not r_nulls.empty:
        merged = pd.concat([merged, l_nulls, r_nulls], ignore_index=True)

    return merged


# 2. Load and merge per TMA file
combined_tmas = {}

for tma, filename in tma_files.items():
    filepath = tma_data_path / filename
    tables = load_excel_tables(filepath)

    print(f"\n{tma}")
    for tbl_name, df in tables.items():
        print(f"  {tbl_name}: {df.shape}")

    table_list = list(tables.values())

    if not table_list:
        print(f"  [Warning] No tables found in {tma}")
        continue

    if len(table_list) == 1:
        combined_tmas[tma] = table_list[0]
    else:
        key = find_common_key(table_list)
        print(f"  -> Merging {len(table_list)} tables using key: '{key}'")
        combined_tmas[tma] = reduce(
            lambda left, right: merge_two_tables(left, right, join_key=key),
            table_list,
        )

    print(f"  -> Combined shape: {combined_tmas[tma].shape}")

# 3. Access separated dataframes
ta4546_combined = combined_tmas.get("TA4546")
ta4547_combined = combined_tmas.get("TA4547")
ta4548_combined = combined_tmas.get("TA4548")

